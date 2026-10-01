#!/usr/bin/env python3
"""svg2pptx.py — turn a finished figure (an SVG that follows the contract of references/spec-guide.md §4)
into an editable PowerPoint deck: one native shape per SVG primitive, grouped per spec element.

  python3 scripts/svg2pptx.py fig.svg                        # → fig.pptx (refuses to overwrite; --force)
  python3 scripts/svg2pptx.py fig.svg -o out.pptx --spec spec.json --scale 2

What becomes what (details in references/pptx-export.md):
  rect / circle / ellipse            → autoshapes (rounded corners kept, fill + outline + dashes exact)
  line / 2-point straight path       → connectors with arrowheads from marker-start / marker-end
  polyline / polygon / other paths   → freeform custom geometry (H V L C S Q T A Z, arcs as béziers)
  <text>                             → text boxes (font mapped from the root stack, CJK companion kept, rotation kept)
  <image> (PNG / JPEG)               → pictures at the same box: data URIs decoded, local files read from the SVG's
                                       folder; preserveAspectRatio kept (meet → fitted, slice → cropped, none → stretched)
  <g data-symbol="tex">              → native PowerPoint equations (OMML) via MathJax MathML + Word's mathml2omml.xsl;
                                       falls back to readable plain text when the stylesheet is not installed
  <g id="n:…|e:…|c:…|l:…|s:…|a:…|legend"> → one PowerPoint group each, named after the id (+ the spec label)
The slide is exactly the viewBox in pt (× --scale), so 1 SVG pt = 1 PowerPoint pt and the layout is identical.
Paint order follows the SVG document order (containers under connectors under boxes); check with pptx_check.py.
"""
import argparse, io, math, os, re, sys
from pathlib import Path
from lxml import etree
sys.path.insert(0, str(Path(__file__).resolve().parent))
import _texkit as T
import _svgkit as K
import measure_text as MT
try:
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
except ImportError:
    sys.exit('python-pptx is required for the PowerPoint export: pip install python-pptx')

SVG = '{%s}' % T.SVG_NS
NS_M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
NS_A14 = 'http://schemas.microsoft.com/office/drawing/2010/main'
NS_MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
VCENTRE = 0.365            # optical centre of a text line = baseline − VCENTRE·size (Latin + CJK sans)
MATH_FONT = 'Cambria Math'
XSL_CANDIDATES = [
    '/Applications/Microsoft Word.app/Contents/Resources/mathml2omml.xsl',
    '/Applications/Microsoft PowerPoint.app/Contents/Resources/mathml2omml.xsl',
    'C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL',
    'C:/Program Files (x86)/Microsoft Office/root/Office16/MML2OMML.XSL',
    'C:/Program Files/Microsoft Office/Office16/MML2OMML.XSL',
    'C:/Program Files/Microsoft Office/Office15/MML2OMML.XSL',
]
# PowerPoint typeface per font key (first installed candidate wins; the last one is the metric-compatible fallback)
PPT_FONTS = {
    'helvetica': ['Helvetica Neue', 'Helvetica', 'Arial'], 'arial': ['Arial'], 'calibri': ['Calibri', 'Carlito', 'Arial'],
    'times': ['Times New Roman', 'Tinos'], 'termes': ['TeX Gyre Termes', 'Times New Roman'],
    'computer-modern': ['Latin Modern Roman', 'CMU Serif', 'Cambria'], 'computer-modern-sans': ['Latin Modern Sans', 'CMU Sans Serif', 'Arial'],
    'dejavu': ['DejaVu Sans', 'Verdana'], 'inter': ['Inter', 'Arial'], 'roboto': ['Roboto', 'Arial'],
    'source-sans': ['Source Sans Pro', 'Source Sans 3', 'Arial'], 'dejavu-mono': ['Menlo', 'Consolas', 'Courier New'],
    'source-han-sans': ['Source Han Sans CN', 'Source Han Sans SC', 'Noto Sans CJK SC', 'PingFang SC', 'Microsoft YaHei'],
    'source-han-serif': ['Source Han Serif CN', 'Source Han Serif SC', 'Noto Serif CJK SC', 'Songti SC', 'SimSun'],
}
WARN = []


def warn(msg):
    if msg not in WARN:
        WARN.append(msg)


# ---------------------------------------------------------------- geometry helpers
def mmul(m, n):                       # m ∘ n: apply n first
    a, b, c, d, e, f = m; A, B, C, D, E, F = n
    return [a * A + c * B, b * A + d * B, a * C + c * D, b * C + d * D, a * E + c * F + e, b * E + d * F + f]


def mapply(m, x, y):
    return (m[0] * x + m[2] * y + m[4], m[1] * x + m[3] * y + m[5])


def parse_transform(s):
    m = [1, 0, 0, 1, 0, 0]
    for name, args in re.findall(r'(\w+)\s*\(([^)]*)\)', s or ''):
        v = [float(t) for t in re.split(r'[\s,]+', args.strip()) if t]
        if name == 'translate': t = [1, 0, 0, 1, v[0], v[1] if len(v) > 1 else 0]
        elif name == 'scale': sx = v[0]; sy = v[1] if len(v) > 1 else sx; t = [sx, 0, 0, sy, 0, 0]
        elif name == 'rotate':
            a = math.radians(v[0]); cx, cy = (v[1], v[2]) if len(v) > 2 else (0, 0)
            ca, sa = math.cos(a), math.sin(a)
            t = mmul([1, 0, 0, 1, cx, cy], mmul([ca, sa, -sa, ca, 0, 0], [1, 0, 0, 1, -cx, -cy]))
        elif name == 'matrix' and len(v) == 6: t = v
        else:
            warn(f'transform {name}() is not supported and was ignored'); continue
        m = mmul(m, t)
    return m


def rot_deg(m):
    return math.degrees(math.atan2(m[1], m[0]))


def scales(m):
    return math.hypot(m[0], m[1]), math.hypot(m[2], m[3])


def parse_color(c, style_opacity=1.0):
    """'#rrggbb' | 'rgb(…)' | named few → (RGBColor, alpha) or (None, 0) for none."""
    c = (c or '').strip()
    if c in ('', 'none', 'transparent'):
        return None, 0.0
    if c.startswith('#'):
        h = c[1:]
        if len(h) == 3: h = ''.join(ch * 2 for ch in h)
        return RGBColor.from_string(h.upper()), style_opacity
    m = re.match(r'rgba?\(([^)]+)\)', c)
    if m:
        v = [float(t.strip().rstrip('%')) for t in m.group(1).split(',')]
        r, g, b = (int(round(x * 2.55)) if '%' in c else int(x) for x in v[:3])
        return RGBColor(r, g, b), (v[3] if len(v) > 3 else 1.0) * style_opacity
    named = {'black': '000000', 'white': 'FFFFFF', 'red': 'FF0000', 'grey': '808080', 'gray': '808080', 'blue': '0000FF', 'green': '008000'}
    if c.lower() in named:
        return RGBColor.from_string(named[c.lower()]), style_opacity
    if c == 'currentColor':
        return RGBColor.from_string('0D0D0D'), style_opacity
    warn(f'colour {c!r} not understood, drawn black'); return RGBColor.from_string('000000'), style_opacity


# ---------------------------------------------------------------- path parsing (full SVG grammar → subpaths)
_NUM = re.compile(r'[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?')


def arc_to_cubics(x1, y1, rx, ry, phi, fa, fs, x2, y2):
    """SVG endpoint arc → list of cubic segments (c1, c2, p) per SVG spec F.6.5."""
    if rx == 0 or ry == 0 or (x1 == x2 and y1 == y2):
        return [((x1, y1), (x2, y2), (x2, y2))]
    phi = math.radians(phi); cp, sp = math.cos(phi), math.sin(phi)
    dx, dy = (x1 - x2) / 2, (y1 - y2) / 2
    x1p, y1p = cp * dx + sp * dy, -sp * dx + cp * dy
    rx, ry = abs(rx), abs(ry)
    lam = (x1p ** 2) / (rx ** 2) + (y1p ** 2) / (ry ** 2)
    if lam > 1: rx *= math.sqrt(lam); ry *= math.sqrt(lam)
    num = rx ** 2 * ry ** 2 - rx ** 2 * y1p ** 2 - ry ** 2 * x1p ** 2
    den = rx ** 2 * y1p ** 2 + ry ** 2 * x1p ** 2
    coef = (1 if fa != fs else -1) * math.sqrt(max(0.0, num / den)) if den else 0.0
    cxp, cyp = coef * rx * y1p / ry, -coef * ry * x1p / rx
    cx, cy = cp * cxp - sp * cyp + (x1 + x2) / 2, sp * cxp + cp * cyp + (y1 + y2) / 2

    def ang(ux, uy, vx, vy):
        d = math.atan2(ux * vy - uy * vx, ux * vx + uy * vy)
        return d
    t1 = ang(1, 0, (x1p - cxp) / rx, (y1p - cyp) / ry)
    dt = ang((x1p - cxp) / rx, (y1p - cyp) / ry, (-x1p - cxp) / rx, (-y1p - cyp) / ry)
    if not fs and dt > 0: dt -= 2 * math.pi
    if fs and dt < 0: dt += 2 * math.pi
    n = max(1, int(math.ceil(abs(dt) / (math.pi / 2) - 1e-9)))
    out, a0 = [], t1
    for i in range(n):
        a1 = t1 + dt * (i + 1) / n
        k = 4 / 3 * math.tan((a1 - a0) / 4)

        def pt(a):
            return (cx + rx * math.cos(a) * cp - ry * math.sin(a) * sp, cy + rx * math.cos(a) * sp + ry * math.sin(a) * cp)

        def dpt(a):
            return (-rx * math.sin(a) * cp - ry * math.cos(a) * sp, -rx * math.sin(a) * sp + ry * math.cos(a) * cp)
        p0, p3 = pt(a0), pt(a1); d0, d3 = dpt(a0), dpt(a1)
        out.append(((p0[0] + k * d0[0], p0[1] + k * d0[1]), (p3[0] - k * d3[0], p3[1] - k * d3[1]), p3))
        a0 = a1
    return out


def parse_path(d):
    """→ list of subpaths: {'pts': [('L', x, y) | ('C', x1, y1, x2, y2, x, y)], 'start': (x, y), 'closed': bool}"""
    toks = re.findall(r'[MmZzLlHhVvCcSsQqTtAa]|' + _NUM.pattern, d or '')
    subs, cur, i = [], None, 0
    x = y = sx = sy = 0.0; lc = None; cmd = None
    while i < len(toks):
        t = toks[i]
        if re.match(r'[A-Za-z]', t):
            cmd = t; i += 1
            if cmd in 'Zz':
                if cur: cur['closed'] = True; x, y = sx, sy
                continue
        else:
            if cmd is None: break
            if cmd == 'M': cmd = 'L'
            elif cmd == 'm': cmd = 'l'

        def num(k):
            return [float(v) for v in toks[i:i + k]]
        if cmd in 'Mm':
            px, py = num(2); i += 2
            if cmd == 'm': px += x; py += y
            x, y = sx, sy = px, py
            cur = {'start': (x, y), 'pts': [], 'closed': False}; subs.append(cur); lc = None
        elif cmd in 'Ll':
            px, py = num(2); i += 2
            if cmd == 'l': px += x; py += y
            x, y = px, py; cur['pts'].append(('L', x, y)); lc = None
        elif cmd in 'Hh':
            px = num(1)[0]; i += 1
            x = px + x if cmd == 'h' else px; cur['pts'].append(('L', x, y)); lc = None
        elif cmd in 'Vv':
            py = num(1)[0]; i += 1
            y = py + y if cmd == 'v' else py; cur['pts'].append(('L', x, y)); lc = None
        elif cmd in 'Cc':
            v = num(6); i += 6
            if cmd == 'c': v = [v[0] + x, v[1] + y, v[2] + x, v[3] + y, v[4] + x, v[5] + y]
            cur['pts'].append(('C', *v)); lc = (v[2], v[3]); x, y = v[4], v[5]
        elif cmd in 'Ss':
            v = num(4); i += 4
            if cmd == 's': v = [v[0] + x, v[1] + y, v[2] + x, v[3] + y]
            c1 = (2 * x - lc[0], 2 * y - lc[1]) if lc else (x, y)
            cur['pts'].append(('C', c1[0], c1[1], v[0], v[1], v[2], v[3])); lc = (v[0], v[1]); x, y = v[2], v[3]
        elif cmd in 'Qq':
            v = num(4); i += 4
            if cmd == 'q': v = [v[0] + x, v[1] + y, v[2] + x, v[3] + y]
            c1 = (x + 2 / 3 * (v[0] - x), y + 2 / 3 * (v[1] - y)); c2 = (v[2] + 2 / 3 * (v[0] - v[2]), v[3] + 2 / 3 * (v[1] - v[3]))
            cur['pts'].append(('C', c1[0], c1[1], c2[0], c2[1], v[2], v[3])); lc = ('Q', v[0], v[1]); x, y = v[2], v[3]
        elif cmd in 'Tt':
            v = num(2); i += 2
            if cmd == 't': v = [v[0] + x, v[1] + y]
            qc = (2 * x - lc[1], 2 * y - lc[2]) if (lc and lc[0] == 'Q') else (x, y)
            c1 = (x + 2 / 3 * (qc[0] - x), y + 2 / 3 * (qc[1] - y)); c2 = (v[0] + 2 / 3 * (qc[0] - v[0]), v[1] + 2 / 3 * (qc[1] - v[1]))
            cur['pts'].append(('C', c1[0], c1[1], c2[0], c2[1], v[0], v[1])); lc = ('Q', qc[0], qc[1]); x, y = v[0], v[1]
        elif cmd in 'Aa':
            v = num(7); i += 7
            ex, ey = (v[5] + x, v[6] + y) if cmd == 'a' else (v[5], v[6])
            for c1, c2, p in arc_to_cubics(x, y, v[0], v[1], v[2], int(v[3]), int(v[4]), ex, ey):
                cur['pts'].append(('C', c1[0], c1[1], c2[0], c2[1], p[0], p[1]))
            x, y = ex, ey; lc = None
        else:
            break
    return subs


def transform_subpaths(subs, m):
    out = []
    for s in subs:
        pts = []
        for p in s['pts']:
            if p[0] == 'L':
                pts.append(('L', *mapply(m, p[1], p[2])))
            else:
                pts.append(('C', *mapply(m, p[1], p[2]), *mapply(m, p[3], p[4]), *mapply(m, p[5], p[6])))
        out.append({'start': mapply(m, *s['start']), 'pts': pts, 'closed': s['closed']})
    return out


# ---------------------------------------------------------------- the converter
class Converter:
    def __init__(self, root, spec, scale, latin, cjk, math_mode, xsl, base_dir=None):
        self.root, self.spec, self.sc = root, spec or {}, scale
        self.base_dir = base_dir                     # relative <image> links resolve against the SVG's folder
        self.W, self.H = self._viewbox()
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Pt(self.W * scale), Pt(self.H * scale)
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.SH = self.slide.shapes
        svg_text = etree.tostring(root, encoding='unicode')
        keys = K.keys_for_svg(svg_text)
        self.fkey = next((k for k in keys if not K.is_cjk_key(k)), None) or K.load_fonts().get('default', 'helvetica')
        self.ckey = next((k for k in keys if K.is_cjk_key(k)), None) or K.cjk_default_key()
        self.latin = latin or self._pick_font(self.fkey)
        self.cjk = cjk or self._pick_font(self.ckey)
        self.root_size = float(root.get('font-size') or 6.5)
        self.math_mode, self.xslt = math_mode, None
        if math_mode == 'omml':
            path = xsl or os.environ.get('MML2OMML_XSL') or next((p for p in XSL_CANDIDATES if os.path.exists(p)), None)
            if path and not os.path.exists(path):
                sys.exit(f'stylesheet not found: {path}')
            if path:
                self.xslt = etree.XSLT(etree.parse(path)); self.xsl_path = path
            else:
                warn('mathml2omml.xsl (ships with Microsoft Word) was not found: formulas are written as plain text; pass --xsl PATH')
        self.mml = {}
        self.markers = self._markers()
        self.made = []          # (gid, label, [shapes]) in document order
        self.counts = {'shape': 0, 'line': 0, 'freeform': 0, 'text': 0, 'math': 0, 'math_text': 0, 'picture': 0}
        self.labels = self._spec_labels()

    def _viewbox(self):
        vb = self.root.get('viewBox')
        if vb:
            x0, y0, w, h = map(float, re.split(r'[\s,]+', vb.strip()))
            if x0 or y0: warn('viewBox does not start at 0 0; coordinates are used as-is')
            return w, h
        return float(re.sub('[a-z]+', '', self.root.get('width') or '240')), float(re.sub('[a-z]+', '', self.root.get('height') or '150'))

    @staticmethod
    def _installed(family):
        pats = [family.replace(' ', ''), family.replace(' ', '-'), family]
        for d in K.load_fonts().get('system_font_dirs', []):
            d = os.path.expanduser(d)
            if not os.path.isdir(d): continue
            try:
                names = os.listdir(d)
            except OSError:
                continue
            for n in names:
                low = n.lower()
                if any(low.startswith(p.lower()) for p in pats):
                    return True
        return False

    def _pick_font(self, key):
        cands = PPT_FONTS.get(key) or [K.load_fonts()['families'].get(key, {}).get('css_family', 'Arial')]
        for c in cands:
            if self._installed(c):
                return c
        return cands[-1]

    def _spec_labels(self):
        lab = {}
        s = self.spec
        for n in s.get('nodes') or []: lab['n:' + n['id']] = n.get('label', '')
        for c in s.get('containers') or []: lab['c:' + c['id']] = c.get('label', '')
        for l in s.get('lanes') or []: lab['l:' + l['id']] = l.get('label', '')
        for a in s.get('annotations') or []: lab['a:' + a['id']] = (a.get('text') or '')[:24]
        for e in s.get('edges') or []: lab['e:' + e['id']] = f"{e.get('from', '')}→{e.get('to', '')}"
        for st in s.get('steps') or []: lab[f"s:{st['n']}"] = f"step {st['n']}"
        lab['legend'] = 'legend'
        return lab

    def _markers(self):
        out = {}
        for m in self.root.iter(SVG + 'marker'):
            paths = [p for p in m.iter() if p.tag in (SVG + 'path', SVG + 'polygon')]
            filled = any((p.get('fill') or 'black') != 'none' for p in paths)
            try: mw = float(m.get('markerWidth') or 3)
            except ValueError: mw = 3.0
            out[m.get('id')] = {'type': 'triangle' if filled else 'arrow', 'mw': mw}
        return out

    # ---- style
    def style_of(self, el, inherited):
        st = dict(inherited)
        for a in ('fill', 'stroke', 'stroke-width', 'stroke-dasharray', 'stroke-linecap', 'stroke-linejoin', 'font-size',
                  'font-weight', 'font-style', 'text-anchor', 'opacity', 'fill-opacity', 'stroke-opacity', 'font-family'):
            v = el.get(a)
            if v not in (None, '', 'inherit'): st[a] = v
        for k, v in re.findall(r'([\w-]+)\s*:\s*([^;]+)', el.get('style') or ''):
            st[k.strip()] = v.strip()
        return st

    def fill_of(self, st):
        return parse_color(st.get('fill', '#000000'), float(st.get('fill-opacity', 1)) * float(st.get('opacity', 1)))

    def stroke_of(self, st):
        return parse_color(st.get('stroke', 'none'), float(st.get('stroke-opacity', 1)) * float(st.get('opacity', 1)))

    # ---- shape styling
    def _alpha(self, clr_el, alpha):
        if alpha < 0.999:
            clr_el.append(clr_el.makeelement(qn('a:alpha'), {'val': str(int(alpha * 100000))}))

    def apply_fill(self, shape, st):
        col, alpha = self.fill_of(st)
        if col is None:
            shape.fill.background()
        else:
            shape.fill.solid(); shape.fill.fore_color.rgb = col
            self._alpha(shape.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr')), alpha)

    def apply_line(self, shape, st, sw_scale=1.0, marker_start=None, marker_end=None):
        col, alpha = self.stroke_of(st)
        ln = shape.line
        if col is None:
            ln.fill.background(); return
        sw = float(st.get('stroke-width', 1)) * sw_scale * self.sc
        ln.color.rgb = col; ln.width = Pt(sw)
        lnel = ln._get_or_add_ln()
        self._alpha(lnel.find(qn('a:solidFill')).find(qn('a:srgbClr')), alpha)
        cap = {'butt': 'flat', 'round': 'rnd', 'square': 'sq'}.get(st.get('stroke-linecap', 'butt'), 'flat')
        lnel.set('cap', cap)
        dash = st.get('stroke-dasharray')
        if dash and dash != 'none':
            ds = [float(v) for v in re.split(r'[\s,]+', dash.strip()) if v]
            if len(ds) % 2: ds = ds * 2
            cd = lnel.makeelement(qn('a:custDash'), {})
            for d, s in zip(ds[::2], ds[1::2]):
                cd.append(cd.makeelement(qn('a:ds'), {'d': str(int(d * self.sc / sw * 100000)), 'sp': str(int(s * self.sc / sw * 100000))}))
            lnel.append(cd)
        join = st.get('stroke-linejoin', 'miter')
        lnel.append(lnel.makeelement(qn('a:' + {'round': 'round', 'bevel': 'bevel'}.get(join, 'miter')), {}))
        for which, mk in (('a:headEnd', marker_start), ('a:tailEnd', marker_end)):
            if mk:
                m = self.markers.get(mk, {'type': 'triangle', 'mw': 3.0})
                size = 'sm' if sw < 0.65 * self.sc else ('lg' if m['mw'] >= 9 else 'med')
                lnel.append(lnel.makeelement(qn(which), {'type': m['type'], 'w': size, 'len': size}))

    @staticmethod
    def no_shadow(shape):
        shape.shadow.inherit = False

    # ---- primitives
    def rect(self, el, m, st, name):
        x, y = float(el.get('x') or 0), float(el.get('y') or 0)
        w, h = float(el.get('width') or 0), float(el.get('height') or 0)
        rx = float(el.get('rx') or el.get('ry') or 0)
        ang = rot_deg(m); sx, sy = scales(m)
        cx, cy = mapply(m, x + w / 2, y + h / 2)
        W, H = w * sx * self.sc, h * sy * self.sc
        shp = self.SH.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rx else MSO_SHAPE.RECTANGLE,
                                Pt(cx * self.sc - W / 2), Pt(cy * self.sc - H / 2), Pt(W), Pt(H))
        if abs(ang) > 0.01: shp.rotation = ang
        if rx:
            gd = shp._element.spPr.find(qn('a:prstGeom')).find(qn('a:avLst'))
            gd.append(gd.makeelement(qn('a:gd'), {'name': 'adj', 'fmla': f'val {int(100000 * min(0.5, rx * sx * self.sc / max(1e-6, min(W, H))))}'}))
        return self._finish(shp, st, name, 'shape')

    def ellipse(self, el, m, st, name):
        cx, cy = float(el.get('cx') or 0), float(el.get('cy') or 0)
        rx = float(el.get('r') or el.get('rx') or 0); ry = float(el.get('r') or el.get('ry') or rx)
        sx, sy = scales(m); X, Y = mapply(m, cx, cy)
        W, H = 2 * rx * sx * self.sc, 2 * ry * sy * self.sc
        shp = self.SH.add_shape(MSO_SHAPE.OVAL, Pt(X * self.sc - W / 2), Pt(Y * self.sc - H / 2), Pt(W), Pt(H))
        ang = rot_deg(m)
        if abs(ang) > 0.01 and abs(rx - ry) > 1e-6: shp.rotation = ang
        return self._finish(shp, st, name, 'shape')

    def _finish(self, shp, st, name, kind):
        shp.name = name; self.no_shadow(shp)
        self.apply_fill(shp, st); self.apply_line(shp, st)
        try: shp.text_frame.word_wrap = False
        except Exception: pass
        self.counts[kind] += 1
        return shp

    def line_like(self, el, m, st, name):
        tag = el.tag.replace(SVG, '')
        if tag == 'line':
            pts = [('L', float(el.get('x1') or 0), float(el.get('y1') or 0)), ('L', float(el.get('x2') or 0), float(el.get('y2') or 0))]
            subs = [{'start': (pts[0][1], pts[0][2]), 'pts': pts[1:], 'closed': False}]
        elif tag in ('polyline', 'polygon'):
            v = [float(t) for t in re.split(r'[\s,]+', (el.get('points') or '').strip()) if t]
            p = list(zip(v[::2], v[1::2]))
            if len(p) < 2: return None
            subs = [{'start': p[0], 'pts': [('L', a, b) for a, b in p[1:]], 'closed': tag == 'polygon'}]
        else:
            subs = parse_path(el.get('d'))
        subs = transform_subpaths([s for s in subs if s['pts']], m)
        if not subs: return None
        ms, me = self._marker_id(el.get('marker-start')), self._marker_id(el.get('marker-end'))
        s0 = subs[0]
        if len(subs) == 1 and len(s0['pts']) == 1 and s0['pts'][0][0] == 'L' and not s0['closed']:
            (x0, y0), (_, x1, y1) = s0['start'], s0['pts'][0]
            shp = self.SH.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x0 * self.sc), Pt(y0 * self.sc), Pt(x1 * self.sc), Pt(y1 * self.sc))
            shp.name = name; self.no_shadow(shp)
            self.apply_line(shp, st, marker_start=ms, marker_end=me)
            self.counts['line'] += 1
            return shp
        return self.freeform(subs, st, name, ms, me)

    @staticmethod
    def _marker_id(v):
        m = re.search(r'#([^)"\']+)', v or '')
        return m.group(1) if m else None

    def freeform(self, subs, st, name, ms, me):
        xs, ys = [], []
        for s in subs:
            xs.append(s['start'][0]); ys.append(s['start'][1])
            for p in s['pts']:
                xs.extend(p[1::2]); ys.extend(p[2::2])
        x0, y0, x1, y1 = min(xs), min(ys), max(xs), max(ys)
        w, h = max(x1 - x0, 0.05), max(y1 - y0, 0.05)
        shp = self.SH.add_shape(MSO_SHAPE.RECTANGLE, Pt(x0 * self.sc), Pt(y0 * self.sc), Pt(w * self.sc), Pt(h * self.sc))
        spPr = shp._element.spPr; prst = spPr.find(qn('a:prstGeom'))
        cust = prst.makeelement(qn('a:custGeom'), {}); prst.addprevious(cust); spPr.remove(prst)
        for t in ('a:avLst', 'a:gdLst', 'a:ahLst', 'a:cxnLst'):
            cust.append(cust.makeelement(qn(t), {}))
        cust.append(cust.makeelement(qn('a:rect'), {'l': '0', 't': '0', 'r': 'r', 'b': 'b'}))
        pl = etree.SubElement(cust, qn('a:pathLst'))
        any_closed = any(s['closed'] for s in subs)
        path = etree.SubElement(pl, qn('a:path'), {'w': str(int(w * 1000)), 'h': str(int(h * 1000))})
        if not any_closed: path.set('fill', 'none')

        def P(parent, x, y):
            parent.append(parent.makeelement(qn('a:pt'), {'x': str(int(round((x - x0) * 1000))), 'y': str(int(round((y - y0) * 1000)))}))
        for s in subs:
            mv = etree.SubElement(path, qn('a:moveTo')); P(mv, *s['start'])
            for p in s['pts']:
                if p[0] == 'L':
                    ln = etree.SubElement(path, qn('a:lnTo')); P(ln, p[1], p[2])
                else:
                    cb = etree.SubElement(path, qn('a:cubicBezTo')); P(cb, p[1], p[2]); P(cb, p[3], p[4]); P(cb, p[5], p[6])
            if s['closed']:
                etree.SubElement(path, qn('a:close'))
        shp.name = name; self.no_shadow(shp)
        if any_closed: self.apply_fill(shp, st)
        else: shp.fill.background()
        self.apply_line(shp, st, marker_start=ms, marker_end=me)
        self.counts['freeform'] += 1
        return shp

    # ---- bitmaps
    def image(self, el, m, st, name):
        """<image> → a native picture. The painted rectangle follows preserveAspectRatio: meet (the default) fits the
        whole picture inside the box, slice fills the box and crops the overflow (PowerPoint's own crop, so it stays
        editable), none stretches it to the box."""
        href = el.get('href') or el.get('{%s}href' % K.XLINK_NS) or ''
        gid = name.rsplit(' ', 1)[0] if ' ' in name else ''; where = gid or 'the figure'
        try:
            data, _ = K.load_image(href, self.base_dir)
        except K.ImageError as ex:
            warn(f'<image> in {where} {ex}: skipped'); return None
        info = K.image_info(data)
        if info['format'] not in K.IMAGE_FORMATS or not info['w']:
            warn(f'<image> in {where} is {info["format"] or "not a bitmap"}, not PNG / JPEG: skipped'); return None
        x, y = float(el.get('x') or 0), float(el.get('y') or 0)
        w, h = (float((K.NUM.match(el.get(a) or '') or [0])[0]) for a in ('width', 'height'))
        if w <= 0 or h <= 0:
            warn(f'<image> in {where} has no width / height: skipped'); return None
        align, mode = K.parse_par(el.get('preserveAspectRatio'))
        bx, by, bw, bh, crop = x, y, w, h, None
        if align != 'none':
            s = (min if mode == 'meet' else max)(w / info['w'], h / info['h'])
            dw, dh = info['w'] * s, info['h'] * s
            fx = {'xMin': 0.0, 'xMid': 0.5, 'xMax': 1.0}[align[:4]]; fy = {'YMin': 0.0, 'YMid': 0.5, 'YMax': 1.0}[align[4:]]
            dx, dy = x + (w - dw) * fx, y + (h - dh) * fy                  # where the whole picture is painted
            if mode == 'meet':
                bx, by, bw, bh = dx, dy, dw, dh
            else:
                crop = ((x - dx) / dw, (y - dy) / dh, (dx + dw - x - w) / dw, (dy + dh - y - h) / dh)
        sx, sy = scales(m); ang = rot_deg(m)
        cx, cy = mapply(m, bx + bw / 2, by + bh / 2)
        W, H = bw * sx * self.sc, bh * sy * self.sc
        pic = self.SH.add_picture(io.BytesIO(data), Pt(cx * self.sc - W / 2), Pt(cy * self.sc - H / 2), Pt(W), Pt(H))
        if crop and any(abs(c) > 1e-6 for c in crop):
            pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
        if abs(ang) > 0.01: pic.rotation = ang
        alpha = float(st.get('opacity', 1))
        if alpha < 0.999:
            blip = pic._element.find('.//' + qn('a:blip'))
            blip.append(blip.makeelement(qn('a:alphaModFix'), {'amt': str(int(alpha * 100000))}))
        pic.name = name                                                 # alt text: the spec label, else the file name
        pic._element.find('.//' + qn('p:cNvPr')).set('descr', self.labels.get(gid) or (os.path.basename(href) if K.href_kind(href) == 'local' else 'bitmap'))
        self.counts['picture'] += 1
        return pic

    # ---- text
    def _run_fonts(self, r, size, bold, italic, col, alpha, math=False):
        r.font.size = Pt(size * self.sc); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = col
        rPr = r._r.get_or_add_rPr()
        self._alpha(rPr.find(qn('a:solidFill')).find(qn('a:srgbClr')), alpha)
        rPr.append(rPr.makeelement(qn('a:latin'), {'typeface': MATH_FONT if math else self.latin}))
        rPr.append(rPr.makeelement(qn('a:ea'), {'typeface': MATH_FONT if math else self.cjk}))

    def _textbox(self, left, top, w, h, anchor, rotation=0.0):
        tb = self.SH.add_textbox(Pt(left), Pt(top), Pt(w), Pt(h))
        tf = tb.text_frame; tf.word_wrap = False; tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = {'start': PP_ALIGN.LEFT, 'middle': PP_ALIGN.CENTER, 'end': PP_ALIGN.RIGHT}.get(anchor, PP_ALIGN.LEFT)
        if abs(rotation) > 0.01: tb.rotation = rotation
        return tb, p

    def text(self, el, m, st, name):
        # SVG (xml:space="default") collapses ASCII whitespace and drops it at both ends; U+3000 etc. stay visible
        s = re.sub(r'[ \t\r\n]+', ' ', ''.join(el.itertext())).strip(' \t\r\n')
        if not s: return None
        size = float(st.get('font-size', self.root_size))
        bold = st.get('font-weight') in ('bold', '600', '700', '800', '900'); italic = st.get('font-style') == 'italic'
        anchor = st.get('text-anchor', 'start')
        x, base = float(el.get('x') or 0), float(el.get('y') or 0)
        col, alpha = parse_color(st.get('fill', '#000000'), float(st.get('fill-opacity', 1)) * float(st.get('opacity', 1)))
        if col is None: col = RGBColor.from_string('0D0D0D')
        w = MT.width_of(s, self.fkey, size, bold=bold, cjk_key=self.ckey)[0] + 2 * size
        h = 2.2 * size
        left = {'start': x, 'middle': x - w / 2, 'end': x - w}.get(anchor, x)
        cy = base - VCENTRE * size
        ang = rot_deg(m)
        # rotate the box centre with the element's transform (gutter labels: rotate(-90 cx cy))
        ccx, ccy = mapply(m, left + w / 2, cy)
        tb, p = self._textbox((ccx - w / 2) * self.sc, (ccy - h / 2) * self.sc, w * self.sc, h * self.sc, anchor, ang)
        tb.name = name
        r = p.add_run(); r.text = s
        self._run_fonts(r, size, bold, italic, col, alpha)
        self.counts['text'] += 1
        return tb

    def math(self, g, m, st, name):
        p = T.params(g, self.root)
        if g.get('data-w') is None:
            r = T.render_svg({'k': {'tex': p['tex'], 'display': p['display']}})['k']
            if 'error' in r:
                warn(f'formula {p["tex"][:40]!r}: {r["error"]}'); return None
            xoff, W, asc, desc = T.metrics(r['svg'], p['size'])
            base = p['y'] + (asc - desc) / 2 if p['valign'] == 'middle' else p['y']
        else:
            W, asc, desc, base = (float(g.get(k)) for k in ('data-w', 'data-asc', 'data-desc', 'data-baseline'))
        size = p['size']
        col, alpha = parse_color(p['fill'], float(st.get('opacity', 1)))
        w = W * 1.35 + 14; h = (asc + desc) * 2.0 + 6
        left = {'start': p['x'], 'middle': p['x'] - w / 2, 'end': p['x'] - w}.get(p['anchor'], p['x'])
        cy = base - (asc - desc) / 2
        ccx, ccy = mapply(m, left + w / 2, cy)
        tb, para = self._textbox((ccx - w / 2) * self.sc, (ccy - h / 2) * self.sc, w * self.sc, h * self.sc, p['anchor'], rot_deg(m))
        tb.name = name
        omml = None
        if self.xslt is not None:
            mml = self.mml.get(p['tex'])
            if mml:
                try:
                    omml = self.xslt(etree.fromstring(mml.encode())).getroot()
                except Exception as ex:
                    warn(f'OMML conversion failed for {p["tex"][:40]!r}: {ex}')
            else:
                warn(f'no MathML for {p["tex"][:40]!r} (MathJax rejected it): written as plain text')
        if omml is None:
            r = para.add_run(); r.text = T.tex_plain(p['tex'])
            self._run_fonts(r, size, False, True, col, alpha, math=True)
            self.counts['math_text'] += 1
            return tb
        for r in omml.iter('{%s}r' % NS_M):
            rPr = r.makeelement(qn('a:rPr'), {'lang': 'en-US', 'sz': str(int(size * self.sc * 100)), 'i': '1', 'dirty': '0'})
            sf = rPr.makeelement(qn('a:solidFill'), {}); c = sf.makeelement(qn('a:srgbClr'), {'val': str(col)}); sf.append(c); rPr.append(sf)
            self._alpha(c, alpha)
            rPr.append(rPr.makeelement(qn('a:latin'), {'typeface': MATH_FONT}))
            mrpr = r.find('{%s}rPr' % NS_M)
            r.insert(1 if mrpr is not None else 0, rPr)
        pel = para._p
        ac = etree.SubElement(pel, '{%s}AlternateContent' % NS_MC, nsmap={'mc': NS_MC})
        ch = etree.SubElement(ac, '{%s}Choice' % NS_MC, nsmap={'a14': NS_A14}); ch.set('Requires', 'a14')
        m14 = etree.SubElement(ch, '{%s}m' % NS_A14)
        opara = etree.SubElement(m14, '{%s}oMathPara' % NS_M, nsmap={'m': NS_M})
        pr = etree.SubElement(opara, '{%s}oMathParaPr' % NS_M)
        etree.SubElement(pr, '{%s}jc' % NS_M).set('{%s}val' % NS_M, {'middle': 'center', 'end': 'right'}.get(p['anchor'], 'left'))
        opara.append(omml)
        fb = etree.SubElement(ac, '{%s}Fallback' % NS_MC)
        r = etree.SubElement(fb, '{%s}r' % NS_A); rp = etree.SubElement(r, '{%s}rPr' % NS_A)
        rp.set('lang', 'en-US'); rp.set('sz', str(int(size * self.sc * 100)))
        etree.SubElement(r, '{%s}t' % NS_A).text = T.tex_plain(p['tex'])
        self.counts['math'] += 1
        return tb

    # ---- walking
    def walk(self, el, m, st, bag, gid):
        for ch in el:
            if not isinstance(ch.tag, str): continue
            tag = ch.tag.replace(SVG, '')
            if tag in ('defs', 'title', 'desc', 'metadata', 'style', 'marker'):
                continue
            cm = mmul(m, parse_transform(ch.get('transform')))
            cst = self.style_of(ch, st)
            label = f'{gid} ' if gid else ''
            if tag == 'g':
                if ch.get('data-symbol') == 'tex':
                    shp = self.math(ch, m, cst, f'{label}公式 {ch.get("data-tex", "")[:24]}')
                    if shp is not None: bag.append(shp)
                else:
                    self.walk(ch, cm, cst, bag, gid)
            elif tag == 'rect': bag.append(self.rect(ch, cm, cst, f'{label}框'))
            elif tag in ('circle', 'ellipse'): bag.append(self.ellipse(ch, cm, cst, f'{label}圆'))
            elif tag in ('path', 'line', 'polyline', 'polygon'):
                shp = self.line_like(ch, cm, cst, f'{label}线')
                if shp is not None: bag.append(shp)
            elif tag == 'text':
                if ch.get('transform'):
                    cm = mmul(m, parse_transform(ch.get('transform')))
                shp = self.text(ch, cm, cst, f'{label}文字')
                if shp is not None: bag.append(shp)
            elif tag == 'image':
                shp = self.image(ch, cm, cst, f'{label}图片')
                if shp is not None: bag.append(shp)
            elif tag in ('use', 'foreignObject', 'switch'):
                warn(f'<{tag}> is not supported and was skipped (the contract forbids it anyway)')
            else:
                warn(f'<{tag}> ignored')

    def run(self, group):
        # formulas: one MathJax call for all of them
        texs = sorted({g.get('data-tex') for g in T.placeholders(self.root) if g.get('data-tex')})
        if texs and self.xslt is not None:
            try:
                res = T.render_mathml({str(i): t for i, t in enumerate(texs)})
                self.mml = {t: res.get(str(i)) for i, t in enumerate(texs)}
            except Exception as ex:
                warn(f'MathJax (node) unavailable, formulas written as plain text: {ex}')
                self.xslt = None
        root_st = self.style_of(self.root, {'fill': '#000000', 'stroke': 'none', 'stroke-width': '1', 'font-size': str(self.root_size)})
        m0 = [1, 0, 0, 1, 0, 0]
        for ch in list(self.root):
            if not isinstance(ch.tag, str): continue
            tag = ch.tag.replace(SVG, '')
            if tag in ('defs', 'title', 'desc', 'metadata', 'style'): continue
            bag = []
            gid = ch.get('id') or ''
            if tag == 'g':
                self.walk(ch, mmul(m0, parse_transform(ch.get('transform'))), self.style_of(ch, root_st), bag, gid)
            else:                                                  # a stray primitive at top level: wrap a copy
                import copy
                tmp = etree.Element(SVG + 'g'); tmp.append(copy.deepcopy(ch))
                self.walk(tmp, m0, root_st, bag, gid)
            self.made.append((gid, bag))
        # group per contract element; legend rows fold into one legend group
        if group:
            legend = [s for gid, bag in self.made if gid.startswith('legend') for s in bag]
            for gid, bag in self.made:
                if gid.startswith('legend'): continue
                if len(bag) > 1:
                    grp = self.SH.add_group_shape(bag)
                    grp.name = f'{gid} {self.labels.get(gid, "")}'.strip() or 'group'
            if len(legend) > 1:
                grp = self.SH.add_group_shape(legend); grp.name = 'legend 图例'
        self._restore_order()

    def _restore_order(self):
        """add_group_shape() appends each new group at the end of the shape tree; put everything back into
        SVG document order so containers stay under connectors and connectors under boxes."""
        spTree = self.SH._spTree
        ext = spTree.find(qn('p:extLst'))

        def top(el):
            while el.getparent() is not None and el.getparent() is not spTree:
                el = el.getparent()
            return el

        def cid(el):
            nv = el.find('.//' + qn('p:cNvPr'))
            return nv.get('id') if nv is not None else None
        want, seen = [], set()
        for gid, bag in self.made:
            for s in bag:
                k = cid(top(s._element))
                if k and k not in seen: seen.add(k); want.append(k)
        index = {cid(c): c for c in spTree if cid(c)}
        for k in want:
            el = index.get(k)
            if el is not None:
                ext.addprevious(el) if ext is not None else spTree.append(el)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('svg')
    ap.add_argument('-o', '--out')
    ap.add_argument('--spec', help='figure spec: group names get the spec labels')
    ap.add_argument('--scale', type=float, default=1.0, help='enlarge everything (2 = twice the print size, easier to edit)')
    ap.add_argument('--latin', help='PowerPoint typeface for Latin text (default: mapped from the SVG font stack)')
    ap.add_argument('--cjk', help='PowerPoint typeface for CJK text (default: mapped from the SVG font stack)')
    ap.add_argument('--math', choices=['omml', 'text'], default='omml', help='formulas as native equations (default) or plain text')
    ap.add_argument('--xsl', help='path to mathml2omml.xsl / MML2OMML.XSL (default: auto-detect from Office)')
    ap.add_argument('--no-group', action='store_true', help='do not group shapes per spec element')
    ap.add_argument('--force', action='store_true', help='overwrite an existing .pptx (it may carry hand edits)')
    a = ap.parse_args()
    out = Path(a.out) if a.out else Path(a.svg).with_suffix('.pptx')
    if out.exists() and not a.force:
        sys.exit(f'{out} exists — it may contain hand edits; pass --force to overwrite or -o for another name')
    spec = None
    if a.spec:
        import json
        spec = json.load(open(a.spec))
    root = T.parse(a.svg).getroot()
    if root.tag != SVG + 'svg':
        sys.exit('not an SVG root element')
    cv = Converter(root, spec, a.scale, a.latin, a.cjk, a.math, a.xsl, base_dir=Path(a.svg).resolve().parent)
    cv.run(group=not a.no_group)
    cv.prs.save(str(out))
    c = cv.counts
    n_groups = 0 if a.no_group else sum(1 for gid, bag in cv.made if len(bag) > 1)
    print(f'wrote {out} — slide {cv.W * a.scale:g}×{cv.H * a.scale:g} pt; {c["shape"]} shapes, {c["line"]} connectors, '
          f'{c["freeform"]} freeforms, {c["text"]} text boxes, {c["math"]} equations' + (f', {c["picture"]} pictures' if c['picture'] else '')
          + (f' (+{c["math_text"]} as plain text)' if c['math_text'] else '') + f'; {n_groups} groups; fonts {cv.latin} / {cv.cjk}'
          + (f'; OMML via {os.path.basename(cv.xsl_path)}' if cv.xslt is not None else ''))
    for w in WARN:
        print('  [warn]', w)


if __name__ == '__main__':
    main()
